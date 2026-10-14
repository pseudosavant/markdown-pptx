from pathlib import Path

import pytest
from pptx import Presentation

from markdown_slides.errors import UnsupportedContentError
from markdown_slides.parser import parse_deck
from markdown_slides.renderer import render_pptx


def parse(source):
    return parse_deck(source, input_path=None, source_name="headings.md")


@pytest.mark.parametrize(
    "deck_layout", [None, "Title and Content", "Title Slide", "Section Header", "Title Only", "Blank", "Two Content"]
)
@pytest.mark.parametrize("heading", ["# Heading", "Heading\n======="])
@pytest.mark.parametrize("subtitle", ["", "## Subtitle", "Subtitle\n--------"])
def test_heading_structure_infers_title_and_section_before_deck_default(deck_layout, heading, subtitle):
    metadata = f"<!-- markdown-pptx:deck\nlayout: {deck_layout}\n-->\n" if deck_layout else ""
    source = f"{metadata}{heading}\n\n{subtitle}\n\n{heading}\n\n{subtitle}\n"
    deck = parse(source)
    assert [slide.layout for slide in deck.slides] == ["Title Slide", "Section Header"]
    for slide in deck.slides:
        assert slide.secondary_body is None
        if subtitle:
            assert slide.body.paragraphs[0].kind == "subtitle"
            assert slide.body.paragraphs[0].heading_level is None
            assert slide.body.paragraphs[0].fragments[0].text == "Subtitle"
        else:
            assert slide.body.is_empty


@pytest.mark.parametrize(
    "body",
    [
        "A paragraph",
        "## First\n\n## Second",
        "## Subtitle\n\nA paragraph",
        "### H3",
        "- Item",
        "> ## Quoted heading",
        "- ## List heading",
        "```markdown\n## Literal heading\n```",
        "![Photo](photo.png)",
        "| A |\n| --- |\n| B |",
        "[Play](demo.mp4)",
    ],
)
def test_other_content_keeps_ordinary_layout_rules(body):
    deck = parse(f"# First\n\n{body}\n\n# Later\n\n{body}\n")
    assert [slide.layout for slide in deck.slides] == ["Title and Content", "Title and Content"]


def test_multiple_h2s_use_configured_deck_default():
    deck = parse("<!-- markdown-pptx:deck\nlayout: Section Header\n-->\n# Heading\n\n## First\n\n## Second\n")
    assert deck.slides[0].layout == "Section Header"
    assert [p.kind for p in deck.slides[0].body.paragraphs] == ["heading", "heading"]


def test_slide_layout_overrides_heading_inference_and_notes_do_not_disable_it():
    deck = parse(
        "# First\n<!-- markdown-pptx:slide\nlayout: Title and Content\n-->\n\n## H2\n"
        "# Section\n<!-- markdown-pptx:slide\nnotes: A section introduction\n-->\n\n## Subtitle\n"
        "# Title only\n<!-- markdown-pptx:slide\nlayout: Title Only\n-->\n"
    )
    assert [s.layout for s in deck.slides] == ["Title and Content", "Section Header", "Title Only"]
    assert deck.slides[0].body.paragraphs[0].kind == "heading"
    assert deck.slides[1].notes == "A section introduction"


def test_dividers_take_precedence_and_empty_heading_stays_blank():
    deck = parse("# Heading\n\n## Left\n\n***\n\n#\n\n# Section\n")
    assert [s.layout for s in deck.slides] == ["Two Content", "Blank", "Section Header"]
    assert deck.slides[0].body.paragraphs[0].kind == "heading"


def test_reference_definitions_and_ordinary_comments_do_not_count_as_content():
    deck = parse(
        "[guide]: https://example.com\n# [Guide][guide]\n\n<!-- Author note -->\n"
        "## [Subtitle][guide]\n\n[other]: https://example.org\n"
    )
    assert deck.slides[0].layout == "Title Slide"
    assert deck.slides[0].body.paragraphs[0].fragments[0].href == "https://example.com"


def test_deck_two_content_still_requires_divider_for_ordinary_content():
    with pytest.raises(UnsupportedContentError, match="exactly one top-level horizontal rule"):
        parse("<!-- markdown-pptx:deck\nlayout: Two Content\n-->\n# Ordinary\n\nA paragraph\n")


def test_native_subtitle_inherits_template_formatting_and_keeps_inline_markup(tmp_path: Path):
    deck = parse("# Title\n\n## Plain **bold** and *italic*\n\n# Section\n\n## Native subtitle\n")
    output = tmp_path / "heading.pptx"
    render_pptx(deck, output_path=output, template_path=None, force=False, base_dir=tmp_path)
    result = Presentation(output)
    assert [s.slide_layout.name for s in result.slides] == ["Title Slide", "Section Header"]
    subtitle = result.slides[0].placeholders[1]
    runs = subtitle.text_frame.paragraphs[0].runs
    assert subtitle.text == "Plain bold and italic"
    assert runs[0].font.bold is None and runs[0].font.italic is None
    assert runs[1].font.bold is True
    assert runs[-1].font.italic is True
    assert all(run.font.size is None for run in runs)
    assert all(not run._r.xpath("./a:rPr/a:latin") for run in runs)
