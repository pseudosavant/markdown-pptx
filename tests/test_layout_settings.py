import itertools
import shutil
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from markdown_slides.errors import ParseError
from markdown_slides.parser import parse_deck
from markdown_slides.renderer import render_pptx


def parse(source):
    return parse_deck(source, input_path=None, source_name="layout.md")


def render(source, tmp_path):
    output = tmp_path / "layout.pptx"
    render_pptx(parse(source), output_path=output, template_path=None, force=True, base_dir=tmp_path)
    return Presentation(output)


def test_layouts_inherit_independently_per_property_and_video_aliases_override_defaults():
    deck = parse(
        "<!-- markdown-pptx:deck\nmedia_layout:\n  width: 75%\n  align: left\n  valign: bottom\n"
        "table_layout:\n  width: 80%\n  align: right\n-->\n"
        "# Comparison\n<!-- markdown-pptx:slide\nmedia_layout:\n  align: right\n"
        "table_layout:\n  align: center\n-->\n\n![Photo](photo.png)\n"
        "<!-- markdown-pptx:media_layout\nvalign: top\n-->\n\n***\n\n| Item |\n| --- |\n| A |\n"
        "<!-- markdown-pptx:table_layout\nwidth: 50%\n-->\n\n"
        "# Video\n\n[Play](demo.mp4)\n<!-- markdown-pptx:video\nmute: true\nalign: center\n-->\n"
        "<!-- markdown-pptx:media_layout\nwidth: auto\n-->\n"
    )
    image = deck.slides[0].body.images[0].layout
    table = deck.slides[0].secondary_body.tables[0].layout
    video = deck.slides[1].body.videos[0]
    assert (image.width, image.align, image.valign) == ("75%", "right", "top")
    assert (table.width, table.align, table.valign) == ("50%", "center", "top")
    assert (video.width, video.align, video.valign, video.mute) == ("auto", "center", "bottom", True)
    assert deck.media_layout.width == "75%"
    assert deck.slides[0].media_layout.specified == {"align"}


def test_empty_layout_mappings_preserve_inherited_placement():
    deck = parse(
        "<!-- markdown-pptx:deck\nmedia_layout:\n  align: left\n  valign: middle\n-->\n"
        "# Photo\n<!-- markdown-pptx:slide\nmedia_layout: {}\n-->\n\n![Photo](photo.png)\n"
        "<!-- markdown-pptx:media_layout\n{}\n-->\n"
    )
    assert deck.slides[0].body.images[0].layout.align == "left"
    assert deck.slides[0].body.images[0].layout.valign == "middle"


@pytest.mark.parametrize("scope", ["deck", "slide", "element"])
@pytest.mark.parametrize("category", ["media_layout", "table_layout"])
@pytest.mark.parametrize(
    "settings",
    [
        "null",
        "[]",
        "true",
        "{width: 0%}",
        "{width: -2in}",
        "{width: 4}",
        "{width: 2px}",
        "{align: null}",
        "{align: []}",
        "{align: true}",
        "{align: justified}",
        "{unknown: top}",
        "{align: left, align: right}",
    ],
)
def test_invalid_layout_settings_fail_at_every_scope(scope, category, settings):
    target = "![Photo](photo.png)" if category == "media_layout" else "| A |\n| --- |\n| B |"
    if scope == "element":
        source = f"# Slide\n\n{target}\n<!-- markdown-pptx:{category}\n{settings}\n-->\n"
    else:
        comment = f"<!-- markdown-pptx:{scope}\n{category}: {settings}\n-->"
        source = f"{comment}\n# Slide\n\n{target}" if scope == "deck" else f"# Slide\n{comment}\n\n{target}"
    with pytest.raises(ParseError) as exc:
        parse(source)
    assert exc.value.context.input_path == "layout.md"
    assert exc.value.context.line is not None


@pytest.mark.parametrize("scope", ["deck", "slide", "element"])
@pytest.mark.parametrize("valign", ["top", "middle", "bottom"])
def test_tables_reject_vertical_alignment_metadata(scope, valign):
    table = "| A |\n| --- |\n| B |"
    if scope == "element":
        source = f"# Table\n\n{table}\n<!-- markdown-pptx:table_layout\nvalign: {valign}\n-->"
    else:
        comment = f"<!-- markdown-pptx:{scope}\ntable_layout:\n  valign: {valign}\n-->"
        source = f"{comment}\n# Table\n\n{table}" if scope == "deck" else f"# Table\n{comment}\n\n{table}"
    with pytest.raises(ParseError, match="Unknown table_layout setting: valign"):
        parse(source)


@pytest.mark.parametrize(
    "body",
    [
        "<!-- markdown-pptx:media_layout\nalign: left\n-->\n\n![Photo](photo.png)",
        "Text\n\n<!-- markdown-pptx:table_layout\nalign: left\n-->",
        "![Photo](photo.png)\n<!-- markdown-pptx:table_layout\nalign: left\n-->",
        "| A |\n| --- |\n| B |\n<!-- markdown-pptx:media_layout\nalign: left\n-->",
        "![Photo](photo.png)\n<!-- markdown-pptx:media_layout\nalign: left\n-->\n<!-- markdown-pptx:media_layout\nwidth: 50%\n-->",
        "[Play](demo.mp4)\n<!-- markdown-pptx:table_layout\nalign: left\n-->",
        "![Photo](photo.png) <!-- markdown-pptx:media_layout align: left -->",
        "| A |\n| --- |\n| <!-- markdown-pptx:table_layout align: left --> |",
        "![Photo](photo.png)\n<!-- markdown-pptx:media_layout align: left -->",
        "![Photo](photo.png)\n<!-- Ordinary comment -->\n<!-- markdown-pptx:media_layout\nalign: left\n-->",
        "- Text\n\n  <!-- markdown-pptx:media_layout\n  align: left\n  -->",
        "> ![Photo](photo.png)\n>\n> <!-- markdown-pptx:media_layout\n> align: left\n> -->",
    ],
)
def test_layout_comments_reject_orphans_wrong_targets_duplicates_and_inline_forms(body):
    with pytest.raises(ParseError):
        parse(f"# Slide\n\n{body}\n")


@pytest.mark.parametrize("gap", ["\n", "\n\n"])
@pytest.mark.parametrize("order", list(itertools.permutations(["video", "media_layout", "media_effects"])))
def test_video_playback_effects_and_layout_comments_can_follow_in_any_order(gap, order):
    comments = {
        "video": "<!-- markdown-pptx:video\nmute: true\nwidth: 75%\n-->",
        "media_layout": "<!-- markdown-pptx:media_layout\nalign: right\nvalign: bottom\n-->",
        "media_effects": "<!-- markdown-pptx:media_effects\nborder_radius: 8pt\n-->",
    }
    video = parse("# Video\n\n[Play](demo.mp4)\n" + gap.join(comments[key] for key in order)).slides[0].body.videos[0]
    assert (video.width, video.align, video.valign, video.mute) == ("75%", "right", "bottom", True)
    assert video.effects.border_radius_pt == 8


@pytest.mark.parametrize("key,value", [("width", "50%"), ("align", "left"), ("valign", "middle")])
@pytest.mark.parametrize("layout_first", [False, True])
def test_duplicate_video_placement_aliases_are_errors(key, value, layout_first):
    comments = [
        f"<!-- markdown-pptx:video\n{key}: {value}\n-->",
        f"<!-- markdown-pptx:media_layout\n{key}: {value}\n-->",
    ]
    if layout_first:
        comments.reverse()
    with pytest.raises(ParseError, match="Duplicate element placement setting"):
        parse("# Video\n\n[Play](demo.mp4)\n" + "\n".join(comments))


@pytest.mark.parametrize("kind", ["image", "local_video", "youtube"])
def test_media_placement_in_both_layouts_uses_each_content_box_and_keeps_native_shapes(kind, tmp_path):
    Image.new("RGB", (160, 90), "blue").save(tmp_path / "photo.png")
    shutil.copy(Path(__file__).parent / "assets/demo.mp4", tmp_path / "demo.mp4")
    target = {
        "image": "![Photo](photo.png)",
        "local_video": "[Play](demo.mp4)",
        "youtube": "[Watch](https://youtu.be/aqz-KE-bpKQ)",
    }[kind]
    cases = list(itertools.product(["left", "center", "right"], ["top", "middle", "bottom"]))
    sources = []
    for index, (align, valign) in enumerate(cases):
        body = ("Context\n\n***\n\n" if index % 2 else "") + target
        sources.append(
            f"# Case {index}\n\n{body}\n<!-- markdown-pptx:media_layout\nwidth: 50%\nalign: {align}\nvalign: {valign}\n-->"
        )
    presentation = render("\n\n".join(sources), tmp_path)
    for index, (align, valign) in enumerate(cases):
        slide = presentation.slides[index]
        shape = next(shape for shape in slide.shapes if shape.name.startswith("MarkdownSlides"))
        area = slide.slide_layout.placeholders[2 if index % 2 else 1]
        assert shape.width == round(area.width * 0.5)
        assert abs(shape.width / shape.height - 16 / 9) < 0.001
        assert shape.left == int(area.left + (area.width - shape.width) * {"left": 0, "center": 0.5, "right": 1}[align])
        assert shape.top == int(
            area.top + (area.height - shape.height) * {"top": 0, "middle": 0.5, "bottom": 1}[valign]
        )
        assert len(slide.shapes) == (3 if index % 2 else 2)
        assert shape._element.xpath(".//p:ph")


def test_default_media_top_alignment_and_portrait_horizontal_override(tmp_path):
    Image.new("RGB", (160, 90), "blue").save(tmp_path / "wide.png")
    Image.new("RGB", (90, 160), "blue").save(tmp_path / "portrait.png")
    source = "# Wide\n\nText\n\n***\n\n![Photo](wide.png)\n# Portrait\n\n![Photo](portrait.png)\n# Left\n\n![Photo](portrait.png)\n<!-- markdown-pptx:media_layout\nalign: left\n-->"
    result = render(source, tmp_path)
    for index, slide in enumerate(result.slides):
        picture = next(s for s in slide.shapes if s.name == "MarkdownSlidesImage")
        area = slide.slide_layout.placeholders[2 if index == 0 else 1]
        assert picture.top == area.top
        expected = area.left if index == 2 else int(area.left + (area.width - picture.width) / 2)
        assert picture.left == expected
        assert picture.width <= area.width and picture.height <= area.height
    wide = next(s for s in result.slides[0].shapes if s.name == "MarkdownSlidesImage")
    assert wide.height < result.slides[0].slide_layout.placeholders[2].height


def test_table_width_and_alignment_keep_zero_height_rows_and_cell_alignment(tmp_path):
    cases = [
        ("auto", "center"),
        ("50%", "left"),
        ("50%", "center"),
        ("50%", "right"),
        ("2in", "right"),
        ("200%", "center"),
    ]
    source = "\n\n".join(
        f"# Table {index}\n\nContext\n\n***\n\n| Item | Value |\n| --- | ---: |\n| A | 1 |\n<!-- markdown-pptx:table_layout\nwidth: {width}\nalign: {align}\n-->"
        for index, (width, align) in enumerate(cases)
    )
    result = render(source, tmp_path)
    for slide, (width, align) in zip(result.slides, cases, strict=True):
        table = next(s for s in slide.shapes if s.has_table)
        area = slide.slide_layout.placeholders[2]
        expected = round(area.width / 2) if width == "50%" else Inches(2) if width == "2in" else area.width
        assert table.width == expected
        assert table.top == area.top
        assert table.left == int(area.left + (area.width - expected) * {"left": 0, "center": 0.5, "right": 1}[align])
        assert [row.height for row in table.table.rows] == [0, 0]
        assert [node.get("anchor") for node in table._element.xpath(".//a:tcPr")] == ["ctr"] * 4
        assert [node.get("algn") for node in table._element.xpath(".//a:pPr")] == ["l", "r", "l", "r"]
        assert len(slide.shapes) == 3


def test_layout_does_not_change_effect_fitting_or_background_geometry(tmp_path):
    Image.new("RGB", (160, 90), "blue").save(tmp_path / "photo.png")
    source = "<!-- markdown-pptx:deck\nbackground: url('photo.png')\nmedia_layout:\n  width: 50%\n  align: right\n  valign: bottom\n-->\n# Plain\n\n![Photo](photo.png)\n# Effects\n\n![Photo](photo.png)\n<!-- markdown-pptx:media_effects\nshadow: 2pt 2pt 6pt\nreflection: 4pt 100%\n-->"
    result = render(source, tmp_path)
    images = [next(s for s in slide.shapes if s.name == "MarkdownSlidesImage") for slide in result.slides]
    assert [(s.left, s.top, s.width, s.height) for s in images] == [
        (images[0].left, images[0].top, images[0].width, images[0].height)
    ] * 2
    assert result.slide_masters[0]._element.xpath("./p:cSld/p:bg/p:bgPr/a:blipFill")
