"""Keep the two product showcases comprehensive and independently renderable."""

import json
import re
from contextlib import closing
from pathlib import Path

import pytest
import yaml
from pptx import Presentation

from markdown_slides.assets import load_syntax_payload
from markdown_slides.effect_settings import MEDIA_EFFECT_KEYS, TABLE_COLOR_KEYS, TABLE_EFFECT_KEYS
from markdown_slides.layout_settings import MEDIA_LAYOUT_KEYS, TABLE_LAYOUT_KEYS
from markdown_slides.markdown_body import MD
from markdown_slides.parser import (
    COLOR_KEYS,
    DOCUMENT_KEYS,
    SLIDE_KEYS,
    SUPPORTED_LAYOUTS,
    TABLE_OPTION_DEFAULTS,
    _split_source,
    parse_deck,
)
from markdown_slides.renderer import Downloader, render_pptx
from markdown_slides.video import VideoDownloader

SAMPLE = Path(__file__).resolve().parents[1] / "sample"
FILES = ("default-showcase.md", "advanced-showcase.md")


def read_showcase(name):
    path = SAMPLE / name
    source = path.read_text(encoding="utf-8")
    document, raw, references = _split_source(source.splitlines(), source_name=name)
    return source, document, raw, references, parse_deck(source, input_path=path, source_name=name)


def element_comments(raw):
    comments = []
    for slide in raw:
        for token in MD.parse(slide.body_markdown):
            if token.type == "html_block":
                for category, values in re.findall(r"<!-- markdown-pptx:(\w+)\n(.*?)\n-->", token.content, re.S):
                    comments.append((category, yaml.safe_load(values)))
    return comments


def test_default_showcase_is_vanilla_markdown_with_native_defaults():
    source, document, raw, _, deck = read_showcase(FILES[0])
    assert "<!-- markdown-pptx:" not in source
    assert document == {}
    assert all(slide.config == {} for slide in raw)
    assert {slide.layout for slide in deck.slides} == {
        "Title Slide",
        "Section Header",
        "Title and Content",
        "Two Content",
        "Blank",
    }
    assert deck.slides[0].layout == "Title Slide"
    for slide in deck.slides:
        assert not slide.hidden
        assert not slide.table_options.total_row
        assert slide.table_options.header_fill is None
        for region in slide.content_regions:
            assert all(item.effects is None for item in [*region.images, *region.tables, *region.videos])
            assert all(item.layout is None for item in [*region.images, *region.tables, *region.videos])


def test_advanced_showcase_covers_deck_and_slide_schema_in_separate_groups():
    _, document, raw, _, deck = read_showcase(FILES[1])
    coverage = json.loads((SAMPLE / "showcase-coverage.json").read_text(encoding="utf-8"))
    groups = coverage["advanced_groups"]
    slides = raw[groups["slide_metadata"] - 1 : groups["element_metadata"] - 1]
    elements = raw[groups["element_metadata"] :]
    assert set(document) == DOCUMENT_KEYS
    assert set(document["fonts"]) == {"body", "headings"}
    assert set(document["color_scheme"]) == COLOR_KEYS | {"preset"}
    assert set(document["media_effects"]) == MEDIA_EFFECT_KEYS
    assert set(document["table_effects"]) == TABLE_EFFECT_KEYS
    assert set(document["table"]) == set(TABLE_OPTION_DEFAULTS) | TABLE_COLOR_KEYS
    assert set(document["media_layout"]) == MEDIA_LAYOUT_KEYS
    assert set(document["table_layout"]) == TABLE_LAYOUT_KEYS
    assert {slide.layout for slide in deck.slides} == SUPPORTED_LAYOUTS
    assert {key for slide in slides for key in slide.config} == SLIDE_KEYS
    assert {key for slide in slides for key in slide.config.get("table", {})} == set(document["table"])
    assert {key for slide in slides for key in slide.config.get("media_effects", {})} == MEDIA_EFFECT_KEYS
    assert {key for slide in slides for key in slide.config.get("table_effects", {})} == TABLE_EFFECT_KEYS
    assert {key for slide in slides for key in slide.config.get("media_layout", {})} == MEDIA_LAYOUT_KEYS
    assert {key for slide in slides for key in slide.config.get("table_layout", {})} == TABLE_LAYOUT_KEYS
    assert element_comments(slides) == []
    assert all(slide.config == {} for slide in elements)
    assert all(element_comments([slide]) for slide in elements)
    comments = element_comments(elements)
    assert {category for category, _ in comments} == {
        "media_effects",
        "table_effects",
        "media_layout",
        "table_layout",
        "video",
    }
    for category, keys in [
        ("media_effects", MEDIA_EFFECT_KEYS),
        ("table_effects", TABLE_EFFECT_KEYS),
        ("media_layout", MEDIA_LAYOUT_KEYS),
        ("table_layout", TABLE_LAYOUT_KEYS),
    ]:
        assert {key for name, settings in comments if name == category for key in settings} == keys
    assert {key for name, settings in comments if name == "video" for key in settings} == {
        "width",
        "aspect_ratio",
        "align",
        "valign",
        "start",
        "fullscreen",
        "loop",
        "mute",
    }
    assert raw[groups["slide_metadata"] - 1].title == "Slide metadata"
    assert raw[groups["element_metadata"] - 1].title == "Element metadata"


def test_showcase_migration_preserves_every_original_topic_without_requiring_old_file():
    coverage = json.loads((SAMPLE / "showcase-coverage.json").read_text(encoding="utf-8"))
    assert coverage["sources"] == list(FILES)
    records = coverage["legacy_migration"]
    assert {record["legacy_slide"] for record in records} == set(range(1, 62))
    sources = {name: read_showcase(name)[2] for name in FILES}
    for record in records:
        slide = sources[record["file"]][record["slide"] - 1]
        assert slide.title_source == record["heading"]
        assert record["legacy_heading"]


def test_showcases_cover_finite_metadata_choices_and_meaningful_combinations():
    _, document, raw, _, deck = read_showcase(FILES[1])
    assert {slide.hidden for slide in deck.slides} == {False, True}
    assert {slide.hide_background_graphics for slide in deck.slides} == {False, True}
    assert {slide.code_highlighting for slide in deck.slides} >= {"default", "theme-dark", "theme-light"}
    assert any(isinstance(slide.master, str) for slide in deck.slides)
    assert any(isinstance(slide.master, int) for slide in deck.slides)
    for key in TABLE_OPTION_DEFAULTS:
        assert {document["table"][key]} | {
            slide.config["table"][key] for slide in raw if key in slide.config.get("table", {})
        } == {False, True}
    assert {slide.background.kind for slide in deck.slides if slide.background} == {
        "none",
        "color",
        "image",
        "gradient",
    }
    assert {
        slide.background.gradient_kind
        for slide in deck.slides
        if slide.background and slide.background.kind == "gradient"
    } == {"linear", "radial"}
    regions = [region for slide in deck.slides for region in slide.content_regions]
    media = [item for region in regions for item in [*region.images, *region.videos]]
    tables = [table for region in regions for table in region.tables]
    for items in [media, tables]:
        assert {item.effects.border.style for item in items if item.effects and item.effects.border} == {
            "solid",
            "dashed",
            "dotted",
        }
        assert any(item.effects and item.effects.border is None for item in items)
        assert any(item.effects and item.effects.shadow and item.effects.reflection for item in items)
    videos = [video for region in regions for video in region.videos]
    assert {video.align for video in videos} == {"left", "center", "right"}
    assert {video.valign for video in videos} == {"top", "middle", "bottom"}
    assert {video.start for video in videos} == {"click", "automatic"}
    assert any(video.width == "auto" for video in videos)
    assert any(video.width.endswith("%") for video in videos)
    assert any(video.width.endswith("in") for video in videos)
    for key in ["fullscreen", "loop", "mute"]:
        assert {getattr(video, key) for video in videos} == {False, True}
    assert {video.kind for video in videos} == {"local_mp4", "remote_mp4", "youtube"}
    assert {item.layout.align for item in media} == {"left", "center", "right"}
    assert {item.layout.valign for item in media} == {"top", "middle", "bottom"}
    assert {item.layout.align for item in tables} == {"left", "center", "right"}
    assert all(item.layout.valign == "top" for item in tables)
    for items in [media, tables]:
        widths = {item.layout.width for item in items}
        assert "auto" in widths
        assert any(width.endswith("%") for width in widths)
        assert any(width.endswith("in") for width in widths)


def walk_fragments(fragments):
    for fragment in fragments:
        yield fragment
        yield from walk_fragments(fragment.children)


def test_showcases_cover_supported_markdown_inventory():
    showcases = [read_showcase(name) for name in FILES]
    source = "\n".join(item[0] for item in showcases)
    slides = [slide for item in showcases for slide in item[4].slides]
    regions = [region for slide in slides for region in slide.content_regions]
    paragraphs = [paragraph for region in regions for paragraph in region.paragraphs]
    fragments = [fragment for paragraph in paragraphs for fragment in walk_fragments(paragraph.fragments)]
    kinds = {fragment.kind for fragment in fragments}
    images = [image for region in regions for image in region.images]
    videos = [video for region in regions for video in region.videos]
    tables = [table for region in regions for table in region.tables]
    tokens = [token for item in showcases for raw in item[2] for token in MD.parse(raw.body_markdown)]
    token_types = {token.type for token in tokens}
    evidence = {
        "paragraphs": any(p.kind == "paragraph" for p in paragraphs),
        "horizontal_rule_as_two_content_separator": any(s.secondary_body is not None for s in slides),
        "bullet_lists": any(p.kind == "list_item" and p.ordered_index is None for p in paragraphs),
        "ordered_lists": any(p.ordered_index is not None for p in paragraphs),
        "nested_lists_up_to_3_levels": max(p.level for p in paragraphs) == 2,
        "list_continuation_paragraphs_and_empty_items": any(p.kind == "list_continuation" for p in paragraphs)
        and any(p.kind == "list_item" and not p.fragments for p in paragraphs),
        "task_lists_with_static_checkboxes": {p.task_checked for p in paragraphs} >= {False, True},
        "atx_or_setext_h1_slide_boundaries": "Setext H1 slide\n===" in source and "# Default showcase" in source,
        "headings_h2_through_h6": {p.heading_level for p in paragraphs} >= set(range(2, 7)),
        "setext_h2_headings": "A smaller heading\n---" in source,
        "emphasis": "emphasis" in kinds,
        "strong": "strong" in kinds,
        "strikethrough": "strike" in kinds,
        "superscript": "superscript" in kinds,
        "subscript": "subscript" in kinds,
        "inline_code": "code" in kinds,
        "links": "link" in kinds,
        "hard_line_breaks": "break" in kinds and "spaces.  \n" in source and "backslash.\\\n" in source,
        "fenced_code_blocks": any(t.type == "fence" and not t.info for t in tokens),
        "fenced_code_syntax_coloring": any(p.kind == "code" and p.code_language for p in paragraphs),
        "indented_code_blocks": "code_block" in token_types,
        "blockquotes": any(p.quote_depth for p in paragraphs),
        "html_tags_and_comments_ignored": "html_block" in token_types and "<span>" in source,
        "pipe_tables": bool(tables),
        "local_images": any(not i.src.startswith("https:") for i in images),
        "remote_images": any(i.src.startswith("https:") for i in images),
        "embedded_h264_mp4_videos": any(v.kind == "local_mp4" for v in videos),
        "online_youtube_videos": any(v.kind == "youtube" for v in videos),
        "standalone_linked_images": any(i.href for i in images),
        "image_alt_text_and_titles": any(i.alt and i.title for i in images),
        "document_wide_link_and_image_references": bool(showcases[0][3])
        and "![local-photo][]" in source
        and "![local-photo]" in source,
        "hyperlink_titles_as_screentips": any(f.kind == "link" and f.title for f in fragments)
        and any(i.link_title for i in images),
        "list_block_indentation": any(p.kind == "code" and p.list_context for p in paragraphs),
        "ordered_list_zero_and_large_starts": {p.ordered_index for p in paragraphs} >= {0, 32768},
        "pipe_table_column_alignment": any(t.alignments == ["left", "center", "right"] for t in tables),
        "native_image_effects": any(i.effects and i.effects.border_radius_pt > 0 for i in images),
        "native_video_effects": any(v.effects and v.effects.border_radius_pt > 0 for v in videos),
        "native_table_effects": any(t.effects and t.effects.border for t in tables),
        "hidden_slides": any(s.hidden for s in slides),
        "heading_based_layout_inference": {s.layout for s in showcases[0][4].slides}
        >= {"Title Slide", "Section Header"},
        "rich_table_cells": {
            fragment.kind
            for table in tables
            for cell in [*table.headers, *(cell for row in table.rows for cell in row)]
            for fragment in walk_fragments(cell)
        }
        >= {"strong", "emphasis", "code", "link", "strike", "subscript", "superscript"},
    }
    assert set(evidence) == set(load_syntax_payload()["supported_markdown"])
    assert all(evidence.values()), [feature for feature, covered in evidence.items() if not covered]


def test_complete_source_alternatives_parse_and_cover_exclusive_deck_choices():
    source = read_showcase(FILES[1])[0]
    snippets = re.findall(r"```markdown\n(.*?)\n```", source, re.S)
    decks = [
        parse_deck(snippet, input_path=SAMPLE / FILES[1], source_name="showcase-example.md") for snippet in snippets
    ]
    assert len(decks) >= 5
    assert {deck.aspect_ratio for deck in decks} == {"4:3", "16:9"}
    assert any(deck.color_scheme and deck.color_scheme.name == "Office" for deck in decks)
    assert {deck.background.kind for deck in decks if deck.background} == {"none", "image"}
    assert any(deck.table_options.banded_columns and not deck.table_options.banded_rows for deck in decks)


@pytest.mark.parametrize("name", FILES)
def test_showcases_render_every_slide_using_bundled_assets(name, tmp_path, monkeypatch):
    deck = read_showcase(name)[4]

    class ShowcaseDownloader(Downloader):
        def fetch(self, url):
            return (SAMPLE / "showcase-local.png").read_bytes()

    monkeypatch.setattr(VideoDownloader, "fetch", lambda self, url: SAMPLE / "big-buck-bunny-clip.mp4")
    output = tmp_path / name.replace(".md", ".pptx")
    template = SAMPLE / "showcase-template.pptx" if name == FILES[1] else None
    with closing(ShowcaseDownloader()) as downloader:
        render_pptx(
            deck, output_path=output, template_path=template, force=False, base_dir=SAMPLE, downloader=downloader
        )
    result = Presentation(output)
    assert len(result.slides) == len(deck.slides)
    assert [slide.slide_layout.name for slide in result.slides] == [slide.layout for slide in deck.slides]
    assert [slide._element.get("show", "1") == "0" for slide in result.slides] == [
        slide.hidden for slide in deck.slides
    ]
