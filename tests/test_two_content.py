from __future__ import annotations

import io
import json
from contextlib import closing
from copy import deepcopy
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from markdown_slides.assets import default_template_path, load_examples
from markdown_slides.cli import main
from markdown_slides.errors import TemplateError, UnsupportedContentError
from markdown_slides.parser import parse_deck
from markdown_slides.renderer import BODY_PLACEHOLDERS, Downloader, list_layout_details, render_pptx


def parse_two(body: str, *, layout: str = "Two Content"):
    return parse_deck(f"# Compare\n---\nlayout: {layout}\n---\n\n{body}\n", input_path=None, source_name="deck.md")


def render(deck, tmp_path: Path, *, template: Path | None = None, downloader=None):
    output = tmp_path / "result.pptx"
    render_pptx(deck, output_path=output, template_path=template, force=True, base_dir=tmp_path, downloader=downloader)
    return Presentation(output)


def content_shapes(slide):
    return sorted(
        [shape for shape in slide.placeholders if shape.placeholder_format.type in BODY_PLACEHOLDERS],
        key=lambda shape: shape.left,
    )


@pytest.mark.parametrize("separator", ["***", "---", "___", "* * *", "-------"])
@pytest.mark.parametrize("layout", ["Two Content", "two-content", "twocontent"])
def test_two_content_commonmark_separators(separator, layout):
    slide = parse_two(f"Left\n\n{separator}\n\nRight", layout=layout).slides[0]
    assert slide.layout == "Two Content"
    assert slide.body.paragraphs[0].fragments[0].text == "Left"
    assert slide.secondary_body.paragraphs[0].fragments[0].text == "Right"


@pytest.mark.parametrize("left,right", [("", "Right"), ("Left", ""), ("", "")])
def test_two_content_empty_sides_render(left, right, tmp_path):
    deck = parse_two(f"{left}\n\n***\n\n{right}")
    slide = render(deck, tmp_path).slides[0]
    assert [shape.text for shape in content_shapes(slide)] == [left, right]
    assert len(slide.shapes) == 3


@pytest.mark.parametrize(
    "body", ["Left", "", "Left\n\n***\n\nMiddle\n\n***\n\nRight", "> ***\n\n***", "- Item\n\n  ***\n\n***"]
)
def test_two_content_rejects_missing_additional_and_nested_breaks(body):
    with pytest.raises(UnsupportedContentError):
        parse_two(body)


def test_separator_error_retains_source_location():
    with pytest.raises(UnsupportedContentError) as exc:
        parse_two("Left\n\n***\n\nRight\n\n***")
    assert exc.value.context.line == 12
    assert exc.value.context.slide_index == 1
    assert exc.value.context.input_path == "deck.md"


def test_fenced_break_and_setext_h2_do_not_divide_content():
    slide = parse_two("Heading\n---\n\n```text\n***\n```\n\n***\n\nRight").slides[0]
    assert slide.body.paragraphs[0].heading_level == 2
    assert slide.body.paragraphs[1].fragments[0].text == "***"
    assert slide.secondary_body.paragraphs[0].fragments[0].text == "Right"


def test_reference_links_resolve_across_regions():
    slide = parse_two("[Left][target]\n\n***\n\n[Right][target]\n\n[target]: https://example.com").slides[0]
    for region in slide.content_regions:
        assert region.paragraphs[0].fragments[0].href == "https://example.com"


@pytest.mark.parametrize("layout", ["Title and Content", "Title Slide", "Section Header", "Title Only", "Blank"])
def test_other_layouts_reject_divider(layout):
    with pytest.raises(UnsupportedContentError, match="layout: Two Content"):
        parse_two("Left\n\n***\n\nRight", layout=layout)


@pytest.mark.parametrize(
    "mixed",
    ["Text\n\n![Image](image.png)", "![A](a.png)\n\n![B](b.png)", "| A |\n| --- |\n| B |\n\n![Image](image.png)"],
)
@pytest.mark.parametrize("side", [0, 1])
def test_each_region_rejects_mixed_or_multiple_objects(mixed, side):
    parts = ["Text", "Text"]
    parts[side] = mixed
    with pytest.raises(UnsupportedContentError):
        parse_two("\n\n***\n\n".join(parts))


def test_two_tables_share_slide_table_options(tmp_path):
    source = "# Tables\n---\nlayout: Two Content\ntable:\n  total_row: true\n  banded_rows: false\n---\n\n"
    source += "| Left |\n| --- |\n| 1 |\n\n***\n\n| Right |\n| --- |\n| 2 |\n"
    deck = parse_deck(source, input_path=None, source_name="deck.md")
    slide = render(deck, tmp_path).slides[0]
    tables = sorted([shape for shape in slide.shapes if shape.has_table], key=lambda shape: shape.left)
    placeholders = content_shapes(slide)
    assert [shape.table.cell(0, 0).text for shape in tables] == ["Left", "Right"]
    for shape, placeholder in zip(tables, placeholders, strict=True):
        assert shape.table.last_row is True
        assert shape.table.horz_banding is False
        assert (shape.left, shape.top, shape.width) == (placeholder.left, placeholder.top, placeholder.width)


def test_text_and_linked_image_use_separate_placeholders(tmp_path):
    Image.new("RGB", (40, 20), "blue").save(tmp_path / "photo.png")
    deck = parse_two('**Left**\n\n***\n\n[![Alternative](photo.png "Picture title")](https://example.com)')
    slide = render(deck, tmp_path).slides[0]
    left, right = content_shapes(slide)
    assert left.text == "Left"
    assert left.text_frame.paragraphs[0].runs[0].font.bold is True
    picture = next(shape for shape in slide.shapes if shape.name == "MarkdownSlidesImage")
    assert picture.left >= right.left
    assert picture.left + picture.width <= right.left + right.width
    assert picture.top >= right.top
    assert picture.top + picture.height <= right.top + right.height
    assert picture.alt_text == "Alternative"
    assert picture.alt_text_title == "Picture title"
    assert picture.click_action.hyperlink.address == "https://example.com"


def test_region_order_follows_position_not_placeholder_index(tmp_path):
    template = Presentation(default_template_path())
    layout = next(item for item in template.slide_layouts if item.name == "Two Content")
    left, right = content_shapes(layout)
    left.left, right.left = right.left, left.left
    path = tmp_path / "reversed.pptx"
    template.save(path)
    slide = render(parse_two("Left\n\n***\n\nRight"), tmp_path, template=path).slides[0]
    assert [shape.text for shape in content_shapes(slide)] == ["Left", "Right"]


@pytest.mark.parametrize("invalid", ["missing", "extra", "overlap"])
def test_layout_inspection_and_render_reject_invalid_content_placeholders(tmp_path, invalid):
    template = Presentation(default_template_path())
    layout = next(item for item in template.slide_layouts if item.name == "Two Content")
    left, right = content_shapes(layout)
    if invalid == "missing":
        right._element.getparent().remove(right._element)
    elif invalid == "extra":
        clone = deepcopy(right._element)
        clone.xpath(".//p:ph")[0].set("idx", "99")
        layout.shapes._spTree.append(clone)
    else:
        right.left = left.left
    path = tmp_path / "invalid.pptx"
    template.save(path)
    detail = next(item for item in list_layout_details(path)["layouts"] if item["name"] == "Two Content")
    assert detail["compatible"] is False
    with pytest.raises(TemplateError) as exc:
        render(parse_two("***"), tmp_path, template=path)
    assert exc.value.context.code == "invalid_content_placeholders"


@pytest.mark.parametrize("topic", [item["id"] for item in load_examples()])
def test_every_example_is_copyable_and_renderable(topic, tmp_path):
    stdout, stderr = io.StringIO(), io.StringIO()
    assert main(["--examples", topic], stdout=stdout, stderr=stderr) == 0
    assert stderr.getvalue() == ""
    source = stdout.getvalue()
    deck = parse_deck(source, input_path=None, source_name=f"{topic}.md")
    Image.new("RGB", (40, 20), "blue").save(tmp_path / "photo.png")

    class ExampleDownloader(Downloader):
        def fetch(self, url):
            return (tmp_path / "photo.png").read_bytes()

    with closing(ExampleDownloader()) as downloader:
        result = render(deck, tmp_path, downloader=downloader)
    assert len(result.slides) == len(deck.slides)


@pytest.mark.parametrize("args", [["--examples"], ["--examples", "list"], ["--examples", "two-content"]])
def test_examples_json(args):
    stdout, stderr = io.StringIO(), io.StringIO()
    assert main([*args, "--json"], stdout=stdout, stderr=stderr) == 0
    payload = json.loads(stdout.getvalue())
    assert payload["mode"] == "examples"
    assert len(payload["examples"]) == (1 if args[-1] == "two-content" else len(load_examples()))
    assert ("markdown" in payload["examples"][0]) == (args[-1] != "list")


@pytest.mark.parametrize(
    "args",
    [
        ["--examples", "unknown"],
        ["--examples", "--syntax"],
        ["--examples", "--force"],
        ["deck.md", "--examples"],
        ["--examples", "--template", "theme.pptx"],
    ],
)
def test_examples_usage_errors(args):
    stdout, stderr = io.StringIO(), io.StringIO()
    assert main([*args, "--json"], stdout=stdout, stderr=stderr) == 2
    assert json.loads(stdout.getvalue())["error"]["code"] == "usage_error"
