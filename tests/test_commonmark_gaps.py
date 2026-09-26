from __future__ import annotations

import pytest
from PIL import Image
from pptx import Presentation

from markdown_slides.errors import ParseError, UnsupportedContentError
from markdown_slides.parser import parse_deck
from markdown_slides.renderer import render_pptx


def parse(source):
    return parse_deck(source, input_path=None, source_name="deck.md")


def render(source, tmp_path):
    path = tmp_path / "deck.pptx"
    render_pptx(parse(source), output_path=path, template_path=None, force=True, base_dir=tmp_path)
    return Presentation(path)


@pytest.mark.parametrize("use", ["[Guide][shared]", "[shared][]", "[shared]"])
@pytest.mark.parametrize("location", ["preamble", "first", "last"])
def test_references_resolve_across_slides_and_titles(use, location):
    definition = '[shared]: https://example.com "Read the guide"\n\n'
    source = f"# {use}\n\n{use}\n\n# Second\n\n{use}\n\n"
    if location == "preamble":
        source = definition + source
    elif location == "first":
        source = source.replace("# Second", definition + "# Second")
    else:
        source += definition
    deck = parse(source)
    fragments = [deck.slides[0].title_fragments[0], *(s.body.paragraphs[0].fragments[0] for s in deck.slides)]
    assert all(f.kind == "link" and f.href == "https://example.com" and f.title == "Read the guide" for f in fragments)


def test_first_definition_wins_with_normalized_labels():
    deck = parse("[A  B]: https://first.example\n\n# [a b]\n\n[A B]\n\n# Last\n\n[a b]: https://last.example\n")
    assert deck.slides[0].title_fragments[0].href == "https://first.example"
    assert deck.slides[0].body.paragraphs[0].fragments[0].href == "https://first.example"


def test_definitions_in_lists_quotes_and_two_content_are_global():
    source = "# [list] [quote] [right]\n\n# Definitions\n\n- [list]: https://list.example\n\n> [quote]: https://quote.example\n\n# Two\n---\nlayout: Two Content\n---\n\nLeft\n\n***\n\n[right]: https://right.example\n"
    deck = parse(source)
    links = [f.href for f in deck.slides[0].title_fragments if f.kind == "link"]
    assert links == ["https://list.example", "https://quote.example", "https://right.example"]


def test_setext_title_reference_and_multiline_definition():
    deck = parse('[Guide][ref]\n===\n\nBody\n\n# Last\n\n[ref]:\n  <https://example.com>\n  "Useful guide"\n')
    assert deck.slides[0].title == "Guide"
    assert deck.slides[0].title_fragments[0].title == "Useful guide"


@pytest.mark.parametrize(
    "definition",
    [
        "```text\n[ref]: https://example.com\n```",
        "<!--\n[ref]: https://example.com\n-->",
        "    [ref]: https://example.com",
    ],
)
def test_literal_definitions_do_not_create_references(definition):
    deck = parse(f"# [ref]\n\n# Literal\n\n{definition}\n")
    assert deck.slides[0].title == "[ref]"


def test_yaml_notes_do_not_define_markdown_references():
    deck = parse("# [ref]\n---\nnotes: |\n  [ref]: https://example.com\n---\n")
    assert deck.slides[0].title == "[ref]"


def test_reference_definition_only_preamble_and_empty_slide():
    deck = parse(
        "---\ncode_highlighting: default\n---\n\n<!-- note -->\n[ref]: https://example.com\n\n#\n\n[other]: https://example.org\n"
    )
    assert deck.slides[0].layout == "Blank"
    with pytest.raises(ParseError, match="Content is not allowed"):
        parse("[ref]: https://example.com\n\nVisible text\n\n# Slide\n")


def test_references_do_not_leak_between_decks():
    parse("[ref]: https://example.com\n\n# [ref]\n")
    assert parse("# [ref]\n").slides[0].title == "[ref]"


def test_ignored_inline_html_can_still_precede_first_slide():
    deck = parse("<span></span>\n\n[ref]: https://example.com\n\n# [ref]\n")
    assert deck.slides[0].title_fragments[0].href == "https://example.com"


def test_referenced_linked_image_preserves_both_titles(tmp_path):
    Image.new("RGB", (30, 20), "blue").save(tmp_path / "picture.png")
    source = '# Picture\n\n[![Alt][picture]][guide]\n\n# Definitions\n\n[picture]: picture.png "Picture title"\n[guide]: https://example.com "Link ScreenTip"\n'
    prs = render(source, tmp_path)
    picture = next(s for s in prs.slides[0].shapes if s.name == "MarkdownSlidesImage")
    assert picture.alt_text == "Alt"
    assert picture.alt_text_title == "Picture title"
    assert picture.click_action.hyperlink.address == "https://example.com"
    assert picture.click_action.hyperlink.screen_tip == "Link ScreenTip"


def test_screen_tips_survive_nested_formatting_in_titles_and_body(tmp_path):
    prs = render(
        '# [**Guide** *title*][ref]\n\n[**Bold** and `code`][ref]\n\n[ref]: https://example.com "Helpful guide"\n',
        tmp_path,
    )
    slide = prs.slides[0]
    runs = [
        r for shape in (slide.shapes.title, slide.placeholders[1]) for p in shape.text_frame.paragraphs for r in p.runs
    ]
    assert all(r.hyperlink.address == "https://example.com" for r in runs)
    assert all(r.hyperlink.screen_tip == "Helpful guide" for r in runs)
    assert any(r.font.bold for r in runs)


@pytest.mark.parametrize("block", ['```python\nprint("Hello")\n```', "## Nested heading", "    indented code"])
def test_blocks_in_list_keep_parent_text_indent(block, tmp_path):
    prs = render("# Slide\n\n1. Run this\n\n   " + block.replace("\n", "\n   ") + "\n", tmp_path)
    paragraphs = prs.slides[0].placeholders[1].text_frame.paragraphs
    assert len(paragraphs) == 2
    assert paragraphs[1].left_indent == paragraphs[0].left_indent
    assert paragraphs[1].first_line_indent == 0


def test_list_quote_nesting_preserves_relative_levels(tmp_path):
    source = "# Slide\n\n- Outer\n\n  > Quoted\n  > - Inner\n  >   - Deep\n  >\n  > Tail\n"
    deck = parse(source)
    paragraphs = deck.slides[0].body.paragraphs
    assert [p.level for p in paragraphs if p.kind == "list_item"] == [0, 1, 2]
    prs = render(source, tmp_path)
    rendered = prs.slides[0].placeholders[1].text_frame.paragraphs
    assert rendered[1].left_indent > rendered[0].left_indent
    assert rendered[3].left_indent > rendered[2].left_indent > rendered[1].left_indent
    assert rendered[4].left_indent == rendered[1].left_indent


def test_quote_does_not_reset_three_level_list_limit():
    with pytest.raises(UnsupportedContentError, match="deeper than 3"):
        parse("# Slide\n\n- Outer\n  > - Inner\n  >   - Deep\n  >     - Too deep\n")


def test_comments_inside_lists_are_ignored_without_extra_items():
    deck = parse("# Slide\n\n- One\n\n  <!-- hidden -->\n\n  continuation\n\n- <!-- empty -->\n- Two\n")
    paragraphs = deck.slides[0].body.paragraphs
    assert [p.kind for p in paragraphs] == ["list_item", "list_continuation", "list_item", "list_item"]
    assert paragraphs[2].fragments == []


@pytest.mark.parametrize("image", ["![alt](pic.png)", "[![alt](pic.png)](https://example.com)", "**![alt](pic.png)**"])
def test_list_images_raise_contextual_error(image):
    with pytest.raises(UnsupportedContentError) as exc:
        parse("# Slide\n\n- " + image + "\n")
    assert "Images are not supported inside list items" in str(exc.value)
    assert exc.value.context.slide_index == 1
    assert exc.value.context.line == 3


@pytest.mark.parametrize("start", [0, 32767, 40000, 999999999])
def test_ordered_numbers_keep_exact_values_without_library_error(start, tmp_path):
    source = f"# Slide\n\n{start}. First\n{start + 1 if start < 999999999 else 1}. Second\n"
    deck = parse(source)
    assert [p.ordered_index for p in deck.slides[0].body.paragraphs] == [start, start + 1]
    prs = render(source, tmp_path)
    for offset, p in enumerate(prs.slides[0].placeholders[1].text_frame.paragraphs):
        number = start + offset
        if 1 <= number <= 32767:
            assert p._p.xpath("./a:pPr/a:buAutoNum")[0].get("startAt") == str(number)
        else:
            assert not p._p.xpath("./a:pPr/a:buAutoNum")
            assert p.text.startswith(f"{number}. ")
