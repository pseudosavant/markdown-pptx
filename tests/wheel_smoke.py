"""Run against an installed package, without pytest's source-path injection."""

from __future__ import annotations

import argparse
import io
import json
import tempfile
from importlib.metadata import version
from pathlib import Path

from pptx import Presentation

from markdown_slides import __version__, skill
from markdown_slides.cli import main


def smoke(*, expect_local: bool) -> None:
    assert version("markdown-pptx") == __version__
    assert skill.is_local_development_build() is expect_local
    original_root = skill.default_skills_dir
    try:
        with tempfile.TemporaryDirectory(prefix="markdown-pptx-wheel-") as directory:
            root = Path(directory) / ".agents" / "skills"
            skill.default_skills_dir = lambda: root
            stdout, stderr = io.StringIO(), io.StringIO()
            assert main(["--version"], stdout=stdout, stderr=stderr) == 0
            assert stdout.getvalue() == f"markdown-pptx {__version__}\n"
            assert stderr.getvalue() == ""
            assert not root.exists()

            # Seed an older skill using the same canonical generator.
            try:
                skill.__version__ = "0"
                assert main(["skill", "install"], stdout=io.StringIO(), stderr=io.StringIO()) == 0
            finally:
                skill.__version__ = __version__
            status = skill.skill_status()
            assert status["managed_version"] == "0"
            assert status["integrity"] == "valid"
            assert status["automatic_sync_eligible"] is not expect_local
            path = Path(status["path"])
            before = path.read_bytes()
            stdout, stderr = io.StringIO(), io.StringIO()
            assert main(["--syntax", "--json"], stdout=stdout, stderr=stderr) == 0
            assert json.loads(stdout.getvalue())["ok"] is True
            if expect_local:
                assert path.read_bytes() == before
                assert stderr.getvalue() == ""
            else:
                assert "Updated managed skill" in stderr.getvalue()
                assert skill.skill_status()["managed_version"] == __version__
            source = Path(directory) / "smoke.md"
            output = Path(directory) / "smoke.pptx"
            source.write_text(
                "<!-- markdown-pptx:deck\nlayout: Title Only\ntable_layout:\n  width: 50%\n"
                "table_effects:\n  shadow: 2pt 2pt 4pt\n-->\n"
                "# Cover\n\n## Native subtitle\n\n# Section\n\n"
                "# Comparison\n\nContext\n\n***\n\n| Item | Value |\n| --- | ---: |\n"
                "| **Bold** | [Guide](https://example.com) |\n"
                "<!-- markdown-pptx:table_layout\nalign: right\n-->\n",
                encoding="utf-8",
            )
            stdout, stderr = io.StringIO(), io.StringIO()
            assert main([str(source), str(output), "--json"], stdout=stdout, stderr=stderr) == 0
            assert json.loads(stdout.getvalue())["slides"] == 3
            presentation = Presentation(output)
            assert [slide.slide_layout.name for slide in presentation.slides] == [
                "Title Slide",
                "Section Header",
                "Two Content",
            ]
            table = next(shape.table for shape in presentation.slides[2].shapes if shape.has_table)
            assert table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.bold is True
            assert table.cell(1, 1).text_frame.paragraphs[0].runs[0].hyperlink.address == "https://example.com"
            # Explicit installation is available for both packaged and local builds.
            assert main(["skill", "install"], stdout=io.StringIO(), stderr=io.StringIO()) == 0
            assert path.read_bytes() == skill.render_skill().encode("utf-8")
            assert "Always invoke the tool as `uvx markdown-pptx ...`" in path.read_text(encoding="utf-8")
            assert list(path.parent.iterdir()) == [path]
            stdout = io.StringIO()
            assert main(["--list-layouts", "--json"], stdout=stdout, stderr=io.StringIO()) == 0
            assert json.loads(stdout.getvalue())["ok"] is True
    finally:
        skill.default_skills_dir = original_root
    print(f"Skill smoke passed (local development: {expect_local})")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--expect-local", action="store_true")
    smoke(expect_local=parser.parse_args().expect_local)
