from __future__ import annotations

import io
import json

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pygments.token import Comment, Keyword, Literal, Name, Operator, Punctuation

from markdown_slides.assets import default_template_path
from markdown_slides.cli import main
from markdown_slides.code_colors import (
    adjusted_rgb,
    contrast,
    fit_color,
    gradient_range,
    image_box,
    image_range,
    luminance,
    solid_range,
    token_role,
)
from markdown_slides.errors import ParseError
from markdown_slides.parser import parse_deck
from markdown_slides.renderer import render_pptx


def source(mode="theme-dark", background="#FFFFFF"):
    return f'---\ncode_highlighting: {mode}\nbackground: "{background}"\n---\n\n# Code\n\n```python\n# readable comment\ndef greet(name: str):\n    return "Hello"\n```\n'


def render(text, tmp_path, template=None):
    deck = parse_deck(text, input_path=None, source_name="deck.md")
    report = {}
    output = tmp_path / "result.pptx"
    render_pptx(deck, output_path=output, template_path=template, force=True, base_dir=tmp_path, report=report)
    return Presentation(output), report


def test_wcag_reference_ratios():
    assert luminance((0, 0, 0)) == 0
    assert luminance((255, 255, 255)) == 1
    assert contrast(0, 1) == 21
    assert contrast(luminance((119, 119, 119)), 1) < 4.5
    assert contrast(luminance((118, 118, 118)), 1) > 4.5


@pytest.mark.parametrize("rgb", [(255, 255, 255), (0, 0, 0), (255, 192, 0), (29, 111, 168), (238, 238, 238)])
@pytest.mark.parametrize("light", [True, False])
def test_minimum_adjustment_meets_target_and_preserves_slot(rgb, light):
    bg = solid_range((0, 0, 0) if light else (255, 255, 255))
    color = fit_color(MSO_THEME_COLOR.ACCENT_1, rgb, bg, light=light)
    assert color.slot == MSO_THEME_COLOR.ACCENT_1
    assert color.ratio >= 4.5
    assert bg.contrast(adjusted_rgb(rgb, color.brightness)) >= 4.5
    if color.brightness:
        previous = color.brightness - (0.00001 if light else -0.00001)
        assert bg.contrast(adjusted_rgb(rgb, previous)) < 4.5


@pytest.mark.parametrize(
    "token,role",
    [
        (Keyword, "keyword"),
        (Keyword.Type, "named"),
        (Keyword.Constant, "value"),
        (Literal.Number, "value"),
        (Literal.String, "value"),
        (Name.Function, "named"),
        (Name.Class, "named"),
        (Name.Builtin, "named"),
        (Name.Tag, "keyword"),
        (Name.Attribute, "named"),
        (Name, "base"),
        (Operator, "base"),
        (Punctuation, "base"),
        (Comment, "comment"),
    ],
)
def test_token_category_mapping(token, role):
    assert token_role(token) == role


@pytest.mark.parametrize("value", ["auto", "light", "THEME-DARK", "null", "false", "[]", "{}"])
@pytest.mark.parametrize("scope", ["document", "slide"])
def test_invalid_metadata(value, scope):
    config = f"---\ncode_highlighting: {value}\n---\n"
    text = config + "\n# Code\n" if scope == "document" else "# Code\n" + config
    with pytest.raises(ParseError) as exc:
        parse_deck(text, input_path=None, source_name="deck.md")
    assert exc.value.context.code == "invalid_code_highlighting"
    assert exc.value.context.input_path == "deck.md"


@pytest.mark.parametrize(
    "mode,background,base",
    [("theme-dark", "#FFFFFF", MSO_THEME_COLOR.DARK_1), ("theme-light", "#102030", MSO_THEME_COLOR.LIGHT_1)],
)
def test_theme_code_is_editable_and_uses_native_theme_references(mode, background, base, tmp_path):
    presentation, report = render(source(mode, background), tmp_path)
    slide = presentation.slides[0]
    paragraph = slide.placeholders[1].text_frame.paragraphs[0]
    runs = {run.text: run for run in paragraph.runs if run.text.strip()}
    assert runs["def"].font.color.theme_color == MSO_THEME_COLOR.ACCENT_1
    assert runs["Hello"].font.color.theme_color == MSO_THEME_COLOR.ACCENT_2
    assert runs["greet"].font.color.theme_color == MSO_THEME_COLOR.ACCENT_3
    assert runs["name"].font.color.theme_color == base
    assert runs["# readable comment"].font.color.theme_color == base
    assert runs["# readable comment"].font.italic
    theme = slide.slide_layout.slide_master.theme.color_scheme
    background_rgb = RGBColor.from_string(background[1:])
    for run in paragraph.runs:
        color = run.font.color
        rgb = adjusted_rgb(theme[color.theme_color], color.brightness)
        assert contrast(luminance(rgb), luminance(background_rgb)) >= 4.5
    assert report["code_highlighting"][0]["target_met"] is True
    assert report["code_highlighting"][0]["warnings"] == []


def test_slide_default_resets_document_theme_mode(tmp_path):
    text = (
        source()
        + '\n# Original palette\n---\ncode_highlighting: default\n---\n\n```python\ndef f():\n    return "Hello"\n```\n'
    )
    p, report = render(text, tmp_path)
    keyword = next(run for run in p.slides[1].placeholders[1].text_frame.paragraphs[0].runs if run.text == "def")
    assert keyword.font.color.rgb == RGBColor.from_string("008000")
    assert len(report["code_highlighting"]) == 1


def test_default_mode_does_not_analyze_background(tmp_path, monkeypatch):
    from markdown_slides.code_background import CodeBackgroundAnalyzer

    monkeypatch.setattr(CodeBackgroundAnalyzer, "background", lambda *args: pytest.fail("default analyzed background"))
    _, report = render(source("default"), tmp_path)
    assert "code_highlighting" not in report


def test_unknown_language_gets_theme_base_color(tmp_path):
    p, _ = render(source().replace("```python", "```unknown-language"), tmp_path)
    assert all(
        r.font.color.theme_color == MSO_THEME_COLOR.DARK_1
        for r in p.slides[0].placeholders[1].text_frame.paragraphs[0].runs
    )


def test_inherited_template_background_is_not_mutated(tmp_path):
    template = Presentation(default_template_path())
    template.slide_master.background.fill.solid()
    template.slide_master.background.fill.fore_color.rgb = RGBColor(20, 30, 40)
    path = tmp_path / "template.pptx"
    template.save(path)
    text = source("theme-light").replace('background: "#FFFFFF"\n', "")
    p, report = render(text, tmp_path, path)
    assert p.slides[0].follow_master_background is True
    assert p.slides[0].background_info.color == RGBColor(20, 30, 40)
    assert report["code_highlighting"][0]["background_method"] == "solid"


def test_two_content_uses_local_gradient_region(tmp_path):
    text = '---\nbackground: "linear-gradient(0deg, #FFFFFF 0%, #E0E0E0 100%)"\ncode_highlighting: theme-dark\n---\n# Split\n---\nlayout: Two Content\n---\n\n```python\ndef a(): pass\n```\n\n***\n\n```python\ndef b(): pass\n```\n'
    _, report = render(text, tmp_path)
    areas = report["code_highlighting"]
    assert len(areas) == 2
    assert all(area["background_method"] == "gradient-grid" and area["samples"] == 256 for area in areas)


def test_image_cropping_and_percentiles_retain_local_brightness():
    image = Image.new("RGB", (200, 100), "white")
    image.paste("black", (100, 0, 200, 100))
    left = image_range(image, image_box(image.size, (0, 0, 100, 100), (0, 0, 50, 100)))
    right = image_range(image, image_box(image.size, (0, 0, 100, 100), (50, 0, 50, 100)))
    assert left.low > 0.99 and right.high < 0.01
    image = Image.new("RGB", (100, 100), "white")
    image.putpixel((50, 50), (0, 0, 0))
    assert image_range(image, (0, 0, 100, 100)).low > 0.99


def test_transparent_image_is_composited():
    image = Image.new("RGBA", (10, 10), (255, 255, 255, 0))
    assert image_range(image, (0, 0, 10, 10), underlay=(0, 0, 0)).high < 0.01


def test_gradient_geometry_and_duplicate_stops():
    stops = [(0, (0, 0, 0)), (0.5, (0, 0, 0)), (0.5, (255, 255, 255)), (1, (255, 255, 255))]
    args = dict(stops=stops, radial=False, center=(0.5, 0.5), slide_size=(100, 100), region=(0, 0, 40, 100))
    assert gradient_range(angle=0, **args).high == 0
    assert gradient_range(angle=180, **args).low == 1
    radial = gradient_range(
        [(0, (255, 255, 255)), (1, (0, 0, 0))],
        angle=0,
        radial=True,
        center=(0.5, 0.5),
        slide_size=(100, 100),
        region=(45, 45, 10, 10),
    )
    assert radial.low > 0.7


def test_impossible_contrast_has_structured_diagnostic(tmp_path):
    _, report = render(source("theme-dark", "#000000"), tmp_path)
    assert report["code_highlighting"][0]["target_met"] is False
    assert report["code_highlighting"][0]["warnings"]


def test_cli_reports_contrast_warning_in_plain_and_json(tmp_path):
    path = tmp_path / "deck.md"
    path.write_text(source("theme-dark", "#000000"), encoding="utf-8")
    for flags in ([], ["--json"]):
        out, err = io.StringIO(), io.StringIO()
        assert main([str(path), "--force", *flags], stdout=out, stderr=err) == 0
        if flags:
            assert json.loads(out.getvalue())["code_highlighting"][0]["warnings"]
            assert err.getvalue() == ""
        else:
            assert "cannot reach 4.5:1" in err.getvalue()


def test_background_image_is_analyzed_once_for_multiple_code_blocks(tmp_path):
    Image.new("RGB", (500, 300), "white").save(tmp_path / "image.png")
    text = source(background="url(image.png)") + "\n```css\n.card { color: blue; }\n```\n"
    _, report = render(text, tmp_path)
    assert len(report["code_highlighting"]) == 1
    assert report["code_highlighting"][0]["background_method"] == "image-percentiles"


def test_unknown_template_background_is_an_assumption(tmp_path):
    template = Presentation(default_template_path())
    template.slide_master.background.fill.patterned()
    path = tmp_path / "pattern.pptx"
    template.save(path)
    _, report = render(source().replace('background: "#FFFFFF"\n', ""), tmp_path, path)
    area = report["code_highlighting"][0]
    assert area["background_method"] == "assumed"
    assert area["target_met"] is None
    assert "assumption" in area["warnings"][0]


@pytest.mark.parametrize("owner", ["layout", "master"])
def test_placeholder_fill_precedes_slide_background(tmp_path, owner):
    from pptx.enum.shapes import PP_PLACEHOLDER

    template = Presentation(default_template_path())
    shape = (
        template.slide_layouts[1].placeholders[1]
        if owner == "layout"
        else template.slide_master.placeholders.get(PP_PLACEHOLDER.BODY)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    path = tmp_path / "filled.pptx"
    template.save(path)
    _, report = render(source(background="#000000"), tmp_path, path)
    assert report["code_highlighting"][0]["target_met"] is True


def test_slide_background_image_cover_uses_only_visible_pixels(tmp_path):
    image = Image.new("RGB", (1000, 100), "black")
    image.paste("white", (400, 0, 600, 100))
    image.save(tmp_path / "wide.png")
    text = '# Code\n---\ncode_highlighting: theme-dark\nbackground: "url(wide.png)"\n---\n\n```python\ndef f(): pass\n```\n'
    _, report = render(text, tmp_path)
    assert report["code_highlighting"][0]["target_met"] is True


def test_image_decoder_is_cached(tmp_path):
    from markdown_slides.code_background import CodeBackgroundAnalyzer

    path = tmp_path / "cached.png"
    Image.new("RGB", (10, 10), "white").save(path)
    analyzer = CodeBackgroundAnalyzer((100, 100))
    try:
        assert analyzer.image(path) is analyzer.image(path)
        assert len(analyzer.images) == 1
    finally:
        analyzer.close()


def test_syntax_describes_theme_modes_and_example():
    output = io.StringIO()
    assert main(["--syntax", "--json"], stdout=output, stderr=io.StringIO()) == 0
    syntax = json.loads(output.getvalue())
    assert syntax["code_highlighting"]["values"] == ["default", "theme-dark", "theme-light"]
    assert syntax["code_highlighting"]["contrast_target"] == 4.5
    assert "code_highlighting" in syntax["document_front_matter_keys"]
    assert "code_highlighting" in syntax["slide_front_matter_keys"]


def test_ignore_document_colors_retains_theme_code_mode(tmp_path):
    input_path = tmp_path / "deck.md"
    input_path.write_text(source(background="#000000"), encoding="utf-8")
    output = io.StringIO()
    assert main([str(input_path), "--json", "--ignore-document-colors"], stdout=output, stderr=io.StringIO()) == 0
    report = json.loads(output.getvalue())
    assert report["code_highlighting"][0]["mode"] == "theme-dark"
    assert report["code_highlighting"][0]["target_met"] is True
