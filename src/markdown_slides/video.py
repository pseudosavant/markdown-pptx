from __future__ import annotations

import io
import tempfile
from pathlib import Path
from urllib.parse import urlsplit

import av
import httpx
from PIL import Image, ImageOps, UnidentifiedImageError
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

from markdown_slides.errors import AssetError
from markdown_slides.models import VideoBlock

DEFAULT_MAX_REMOTE_VIDEO_MB = 100


class VideoDownloader:
    def __init__(
        self, *, enabled: bool = True, max_mb: int = DEFAULT_MAX_REMOTE_VIDEO_MB, client: httpx.Client | None = None
    ) -> None:
        self.enabled = enabled
        self.max_bytes = max_mb * 1_000_000
        self._client = client or httpx.Client(follow_redirects=True, timeout=httpx.Timeout(30.0, read=120.0))
        self._owns_client = client is None
        self._cache: dict[str, Path] = {}

    def fetch(self, url: str) -> Path:
        host = urlsplit(url).hostname or "remote host"
        if not self.enabled:
            raise AssetError("remote_videos_disabled", f"Remote videos are disabled: {host}")
        if url in self._cache:
            return self._cache[url]
        path: Path | None = None
        try:
            with self._client.stream("GET", url) as response:
                response.raise_for_status()
                if response.url.scheme != "https":
                    raise AssetError("invalid_remote_video_url", f"Remote video redirected outside HTTPS: {host}")
                content_type = response.headers.get("content-type", "").split(";", 1)[0].strip().lower()
                if content_type and content_type not in {
                    "video/mp4",
                    "application/octet-stream",
                    "binary/octet-stream",
                }:
                    raise AssetError("invalid_remote_video_type", f"Remote video is not MP4 ({content_type}): {host}")
                length = response.headers.get("content-length")
                if length and length.isdecimal() and int(length) > self.max_bytes:
                    raise AssetError(
                        "remote_video_too_large",
                        f"Remote video exceeds the {self.max_bytes // 1_000_000} MB limit: {host}",
                    )
                with tempfile.NamedTemporaryFile(delete=False, suffix=".mp4") as target:
                    path = Path(target.name)
                    total = 0
                    for chunk in response.iter_bytes():
                        total += len(chunk)
                        if total > self.max_bytes:
                            raise AssetError(
                                "remote_video_too_large",
                                f"Remote video exceeds the {self.max_bytes // 1_000_000} MB limit: {host}",
                            )
                        target.write(chunk)
            self._cache[url] = path
            return path
        except AssetError:
            raise
        except httpx.HTTPError as exc:
            raise AssetError(
                "video_download_failed", f"Could not download remote video from {host}: {type(exc).__name__}"
            ) from exc
        except OSError as exc:
            raise AssetError("video_download_failed", f"Could not save remote video from {host}: {exc}") from exc
        finally:
            if path is not None and url not in self._cache:
                path.unlink(missing_ok=True)

    def close(self) -> None:
        if self._owns_client:
            self._client.close()
        for path in self._cache.values():
            path.unlink(missing_ok=True)
        self._cache.clear()


def render_video(slide, placeholder, video: VideoBlock, *, base_dir: Path, downloader: VideoDownloader) -> None:
    if video.kind == "youtube":
        aspect = video.aspect_ratio or 16 / 9
        left, top, width, height = _placement(placeholder, video, aspect)
        picture = slide.shapes.add_picture(io.BytesIO(_youtube_cover()), left, top, width, height)
        picture.name = "MarkdownSlidesYouTubeVideo"
        relationship = slide.part.relate_to(video.source, RT.VIDEO, is_external=True)
        _attach_video_link(picture, relationship)
        _add_media_timing(slide, picture.shape_id)
        return

    path = downloader.fetch(video.source) if video.kind == "remote_mp4" else (base_dir / video.source).resolve()
    if not path.is_file():
        raise AssetError("missing_asset", f"Video asset does not exist: {path}")
    aspect, first_frame = _inspect_video(path)
    aspect = video.aspect_ratio or aspect
    poster = _poster(video.poster, first_frame, aspect, base_dir)
    left, top, width, height = _placement(placeholder, video, aspect)
    shape = slide.shapes.add_movie(
        str(path), left, top, width, height, poster_frame_image=poster, mime_type="video/mp4"
    )
    shape.name = "MarkdownSlidesVideo"
    _configure_local_video(slide, shape, video)


def _inspect_video(path: Path) -> tuple[float, Image.Image]:
    try:
        with av.open(str(path)) as container:
            if "mp4" not in container.format.name.split(","):
                raise AssetError("invalid_video", f"Video is not an MP4 file: {path}")
            stream = next((item for item in container.streams.video), None)
            if stream is None or stream.codec_context.name != "h264":
                raise AssetError("invalid_video_codec", f"Video needs an H.264 stream: {path}")
            ratio = stream.sample_aspect_ratio or 1
            aspect = float(stream.width * ratio / stream.height)
            if aspect <= 0:
                raise AssetError("invalid_video", f"Video has invalid dimensions: {path}")
            frame = next(container.decode(stream))
            return aspect, frame.to_image()
    except AssetError:
        raise
    except (OSError, ValueError, StopIteration, ZeroDivisionError, av.FFmpegError) as exc:
        raise AssetError("invalid_video", f"Could not decode H.264 MP4 video: {path} ({type(exc).__name__})") from exc


def _poster(poster_path: str | None, first_frame: Image.Image, aspect: float, base_dir: Path) -> io.BytesIO:
    if poster_path is None:
        image = first_frame
    else:
        path = (base_dir / poster_path).resolve()
        if not path.is_file():
            raise AssetError("missing_asset", f"Video poster does not exist: {path}")
        try:
            with Image.open(path) as source:
                source.load()
                image = source.convert("RGB")
        except (OSError, UnidentifiedImageError) as exc:
            raise AssetError("invalid_image", f"Video poster could not be decoded: {path}") from exc
    target_width = min(max(image.width, 320), 1920)
    target_height = max(1, round(target_width / aspect))
    if target_height > 1080:
        target_height = 1080
        target_width = max(1, round(target_height * aspect))
    fitted = ImageOps.contain(image.convert("RGB"), (target_width, target_height))
    canvas = Image.new("RGB", (target_width, target_height), "black")
    canvas.paste(fitted, ((target_width - fitted.width) // 2, (target_height - fitted.height) // 2))
    result = io.BytesIO()
    canvas.save(result, format="PNG")
    result.seek(0)
    return result


def _placement(placeholder, video: VideoBlock, aspect: float) -> tuple[int, int, int, int]:
    max_width = int(placeholder.width)
    max_height = int(placeholder.height)
    if video.width.endswith("%"):
        desired = round(max_width * float(video.width[:-1]) / 100)
    elif video.width.endswith("in"):
        desired = int(Inches(float(video.width[:-2])))
    else:
        desired = max_width
    width = min(max_width, desired, round(max_height * aspect))
    height = round(width / aspect)
    horizontal = {"left": 0, "center": 0.5, "right": 1}[video.align]
    vertical = {"top": 0, "middle": 0.5, "bottom": 1}[video.valign]
    left = int(placeholder.left + (max_width - width) * horizontal)
    top = int(placeholder.top + (max_height - height) * vertical)
    return left, top, width, height


def _element(tag: str, **attributes):
    element = OxmlElement(tag)
    for key, value in attributes.items():
        element.set(key, str(value))
    return element


def _child(parent, tag: str, **attributes):
    element = _element(tag, **attributes)
    parent.append(element)
    return element


def _add_media_timing(slide, shape_id: int):
    timing = slide._element.find(qn("p:timing"))
    if timing is None:
        timing = _child(slide._element, "p:timing")
        timeline = _child(_child(timing, "p:tnLst"), "p:par")
        root = _child(timeline, "p:cTn", id="1", dur="indefinite", restart="never", nodeType="tmRoot")
        _child(root, "p:childTnLst")
    root_children = timing.find(".//" + qn("p:cTn") + "[@nodeType='tmRoot']/" + qn("p:childTnLst"))
    next_id = max(int(node.get("id")) for node in timing.iter(qn("p:cTn"))) + 1
    media = _child(root_children, "p:video")
    media_node = _child(media, "p:cMediaNode", vol="80000")
    media_time = _child(media_node, "p:cTn", id=next_id, fill="hold", display="0")
    _child(_child(media_time, "p:stCondLst"), "p:cond", delay="indefinite")
    _child(_child(media_node, "p:tgtEl"), "p:spTgt", spid=shape_id)
    return media, media_node, media_time


def _attach_video_link(picture, relationship: str) -> None:
    nv_pic = picture._element.find(qn("p:nvPicPr"))
    properties = nv_pic.find(qn("p:cNvPr"))
    _child(properties, "a:hlinkClick", action="ppaction://media").set(qn("r:id"), "")
    nv_pr = nv_pic.find(qn("p:nvPr"))
    video_file = _child(nv_pr, "a:videoFile")
    video_file.set(qn("r:link"), relationship)


def _configure_local_video(slide, shape, video: VideoBlock) -> None:
    timing = slide._element.find(qn("p:timing"))
    root_children = timing.find(".//" + qn("p:cTn") + "[@nodeType='tmRoot']/" + qn("p:childTnLst"))
    media = next(
        item
        for item in root_children.findall(qn("p:video"))
        if item.find(".//" + qn("p:spTgt")).get("spid") == str(shape.shape_id)
    )
    media_node = media.find(qn("p:cMediaNode"))
    media_node.set("vol", "0" if video.mute else "80000")
    media_time = media_node.find(qn("p:cTn"))
    if video.loop:
        media_time.set("repeatCount", "indefinite")
    if video.fullscreen:
        media.set("fullScrn", "1")
    sequence = root_children.find(qn("p:seq"))
    next_id = max(int(node.get("id")) for node in timing.iter(qn("p:cTn"))) + 1
    if sequence is None:
        sequence = _element("p:seq", concurrent="1", nextAc="seek")
        root_children.insert(0, sequence)
        main = _child(sequence, "p:cTn", id=next_id, dur="indefinite", nodeType="mainSeq")
        _child(main, "p:childTnLst")
        for direction in ("prev", "next"):
            condition = _child(
                _child(sequence, f"p:{direction}CondLst"),
                "p:cond",
                evt="onPrev" if direction == "prev" else "onNext",
                delay="0",
            )
            _child(_child(condition, "p:tgtEl"), "p:sldTgt")
        next_id += 1
    main = sequence.find(qn("p:cTn"))
    main_id = main.get("id")
    outer = _child(main.find(qn("p:childTnLst")), "p:par")
    outer_time = _child(outer, "p:cTn", id=next_id, fill="hold")
    conditions = _child(outer_time, "p:stCondLst")
    _child(conditions, "p:cond", delay="indefinite")
    if video.start == "automatic":
        begin = _child(conditions, "p:cond", evt="onBegin", delay="0")
        _child(begin, "p:tn", val=main_id)
    middle = _child(_child(outer_time, "p:childTnLst"), "p:par")
    middle_time = _child(middle, "p:cTn", id=next_id + 1, fill="hold")
    _child(_child(middle_time, "p:stCondLst"), "p:cond", delay="0")
    inner = _child(_child(middle_time, "p:childTnLst"), "p:par")
    effect = _child(
        inner,
        "p:cTn",
        id=next_id + 2,
        presetID="1",
        presetClass="mediacall",
        presetSubtype="0",
        fill="hold",
        nodeType="withEffect" if video.start == "automatic" else "clickEffect",
    )
    _child(_child(effect, "p:stCondLst"), "p:cond", delay="0")
    command = _child(_child(effect, "p:childTnLst"), "p:cmd", type="call", cmd="playFrom(0.0)")
    behavior = _child(command, "p:cBhvr")
    _child(behavior, "p:cTn", id=next_id + 3, dur="1", fill="hold")
    _child(_child(behavior, "p:tgtEl"), "p:spTgt", spid=shape.shape_id)


def _youtube_cover() -> bytes:
    image = Image.new("RGB", (640, 360), "#20242b")
    from PIL import ImageDraw

    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((266, 138, 374, 222), radius=21, fill="#ef2428")
    draw.polygon([(308, 159), (308, 201), (344, 180)], fill="white")
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()
