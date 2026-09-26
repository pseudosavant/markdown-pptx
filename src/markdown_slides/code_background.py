"""Connect background analysis to public PowerPoint and image APIs."""

from __future__ import annotations

import io
import warnings
from pathlib import Path

from PIL import Image
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL, MSO_THEME_COLOR
from pptx.enum.shapes import PP_PLACEHOLDER

from markdown_slides.code_colors import adjusted_rgb, gradient_range, image_box, image_range, make_palette, solid_range
from markdown_slides.errors import AssetError


class CodeBackgroundAnalyzer:
    def __init__(self, slide_size):
        self.slide_size = slide_size
        self.images = {}
        self.reports = []

    def close(self):
        for image in self.images.values():
            image.close()

    def image(self, source):
        key = source.getvalue() if isinstance(source, io.BytesIO) else str(Path(source).resolve())
        if key not in self.images:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                try:
                    if isinstance(source, io.BytesIO):
                        source.seek(0)
                    image = Image.open(source)
                    if image.width * image.height > 50_000_000:
                        image.close()
                        raise AssetError("image_dimensions_too_large", "Background image exceeds the pixel limit.")
                    image.load()
                    self.images[key] = image
                except (OSError, Image.DecompressionBombError, Image.DecompressionBombWarning) as exc:
                    raise AssetError("invalid_image", "Background image could not be analyzed.") from exc
        return self.images[key]

    def palette(self, slide, spec, deck, placeholder, body, *, resolve_image, fit_image):
        mode = spec.code_highlighting or deck.code_highlighting
        if mode == "default" or not any(p.kind == "code" for p in body.paragraphs):
            return None
        scheme = slide.slide_layout.slide_master.theme.color_scheme
        region = (placeholder.left, placeholder.top, placeholder.width, placeholder.height)
        notices = []
        assumed = False
        try:
            background = self.background(slide, spec, deck, placeholder, region, scheme, resolve_image, fit_image)
        except (ValueError, KeyError, AttributeError, TypeError) as exc:
            assumed = True
            opposite = MSO_THEME_COLOR.DARK_1 if mode == "theme-light" else MSO_THEME_COLOR.LIGHT_1
            background = solid_range(scheme[opposite])
            notices.append(
                f"Background could not be analyzed ({exc}). Contrast uses the opposite theme text color as an assumption."
            )
        palette = make_palette(scheme, mode, background)
        minimum = min(color.ratio for color in palette.values())
        if minimum < 4.5:
            notices.append(
                "The selected light/dark variant cannot reach 4.5:1 over the analyzed background. The requested variant is retained at its maximum contrast."
            )
        self.reports.append(
            {
                "slide": spec.index,
                "placeholder": placeholder.placeholder_format.idx,
                "mode": mode,
                "background_method": "assumed" if assumed else background.method,
                "samples": background.samples,
                "minimum_contrast": round(minimum, 4),
                "target_met": None if assumed else minimum >= 4.5,
                "warnings": notices,
            }
        )
        return palette

    def background(self, slide, spec, deck, placeholder, region, scheme, resolve_image, fit_image):
        # Inspect explicit placeholder fills before the slide background. No floating text boxes.
        shapes = [placeholder]
        layout_shape = next(
            (
                p
                for p in slide.slide_layout.placeholders
                if p.placeholder_format.idx == placeholder.placeholder_format.idx
            ),
            None,
        )
        if layout_shape is not None:
            shapes.append(layout_shape)
        # Code is only rendered in body, object, or subtitle placeholders.
        # All three inherit the body placeholder on the master.
        master_shape = slide.slide_layout.slide_master.placeholders.get(PP_PLACEHOLDER.BODY, None)
        if master_shape is not None:
            shapes.append(master_shape)
        for shape in shapes:
            fill = shape.fill
            if fill.type == MSO_FILL.SOLID:
                color = fill.fore_color
                rgb = color.rgb if color.type == MSO_COLOR_TYPE.RGB else scheme[color.theme_color]
                return solid_range(adjusted_rgb(rgb, color.brightness))
            if fill.type == MSO_FILL.BACKGROUND:
                break
            if fill.type is not None:
                raise ValueError("unsupported content placeholder fill")
        info = slide.background_info
        # Slide image backgrounds are pictures behind content. Document images are
        # stretched master fills and are already represented by background_info.
        if spec.background is not None and spec.background.kind == "image":
            image_source = resolve_image(spec.background.url or "")
            image = self.image(image_source)
            placement = fit_image(image.size, 0, 0, *self.slide_size, contain=False)
            underlay = info.color if info.kind == "solid" else (255, 255, 255)
            if ("A" in image.getbands() or "transparency" in image.info) and info.kind != "solid":
                raise ValueError("transparent image over a non-solid background")
            return image_range(image, image_box(image.size, placement, region), underlay=underlay)
        if info.kind == "solid":
            return solid_range(info.color)
        if info.kind == "gradient":
            return gradient_range(
                info.stops,
                angle=info.angle,
                radial=info.radial,
                center=info.center,
                slide_size=self.slide_size,
                region=region,
            )
        if info.kind == "picture":
            image = self.image(io.BytesIO(info.image))
            # Native background fills are composited onto the presentation canvas.
            underlay = (255, 255, 255)
            left, top, right, bottom = info.crop
            box = image_box(image.size, (0, 0, *self.slide_size), region)
            crop_box = (
                left * image.width + box[0] * (1 - left - right),
                top * image.height + box[1] * (1 - top - bottom),
                left * image.width + box[2] * (1 - left - right),
                top * image.height + box[3] * (1 - top - bottom),
            )
            return image_range(image, crop_box, underlay=underlay)
        if info.kind == "none":
            return solid_range((255, 255, 255))
        raise ValueError(info.reason or "unknown template background")
